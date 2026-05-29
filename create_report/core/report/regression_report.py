from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from create_report.core.report.base_report import BaseReport
from create_report.core.settings_reader import get_settings
from version_detector import find_version


class RegressionReport(BaseReport):

    def generate(self, datasets: list, output_path: str, country: str, version: str):
        wb = Workbook()
        ws = wb.active
        ws.title = "Регресс"

        # =========================
        # ЗАГОЛОВОК
        # =========================
        ws["E1"] = f"Регресс версии {version}"
        ws["E1"].font = Font(bold=True, size=14)

        # =========================
        # ШАПКА
        # =========================
        start_row = 3

        ws.cell(row=start_row, column=1, value="Регион")
        ws.cell(row=start_row + 1, column=1, value=country)

        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=4)
        ws.cell(row=start_row, column=2, value="Версия ПО")

        ws.cell(row=start_row + 1, column=2, value="Securos")
        ws.cell(row=start_row + 1, column=3, value="Auto")
        ws.cell(row=start_row + 1, column=4, value="NN-server")

        ws.merge_cells(start_row=start_row, start_column=5, end_row=start_row + 1, end_column=5)
        ws.cell(row=start_row, column=5, value="Пайплайн")

        ws.merge_cells(start_row=start_row, start_column=6, end_row=start_row, end_column=7)
        ws.cell(row=start_row, column=6, value="Качество распознавания")
        ws.cell(row=start_row + 1, column=6, value="Полное совпадение (%)")
        ws.cell(row=start_row + 1, column=7, value="Включая ошибочные (%)")

        ws.merge_cells(start_row=start_row, start_column=8, end_row=start_row, end_column=9)
        ws.cell(row=start_row, column=8, value="Качество детекции")
        ws.cell(row=start_row + 1, column=8, value="Пропущено (%)")
        ws.cell(row=start_row + 1, column=9, value="Повторов (%)")

        ws.merge_cells(start_row=start_row, start_column=10, end_row=start_row, end_column=11)
        ws.cell(row=start_row, column=10, value="Изменения относительно прошлой версии")
        ws.cell(row=start_row + 1, column=10, value="Качество распознавания (%)")
        ws.cell(row=start_row + 1, column=11, value="Качество детекции (%)")

        # =========================
        # ДАННЫЕ
        # =========================
        row = start_row + 2

        for dataset in datasets:
            name = dataset["name"]
            stats = dataset["stats"]
            scenario_name = f"{country} - {name}"

            ws.cell(row=row, column=1, value=scenario_name)
            ws.cell(row=row, column=2, value=find_version("SecurOS"))
            ws.cell(row=row, column=3, value=find_version("SecurOS Auto"))

            nn_value = (
                "Auto - " + find_version("SecurOS NN-Server Auto for Intel") +
                "\nVehicle - " + find_version("SecurOS NN-Server Vehicle for Intel")
            )
            cell = ws.cell(row=row, column=4, value=nn_value)
            cell.alignment = Alignment(wrap_text=True)

            ws.cell(row=row, column=5, value="NN")
            ws.cell(row=row, column=6, value=round(stats["f1_strict"], 2))
            ws.cell(row=row, column=7, value=round(stats["f1_soft"], 2))
            ws.cell(row=row, column=8, value=round(stats["missed_count"] / stats["etalon_total"] * 100, 2) if stats["etalon_total"] else 0)
            ws.cell(row=row, column=9, value=round(stats["duplicates"] / stats["total"] * 100, 2) if stats["total"] else 0)

            row += 1

        last_data_row = row - 1

        # =========================
        # СТИЛИ ЛИСТА РЕГРЕСС
        # =========================
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for row_cells in ws.iter_rows(min_row=start_row, max_row=last_data_row, min_col=1, max_col=11):
            for cell in row_cells:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

        for r in range(start_row, last_data_row + 1):
            ws.row_dimensions[r].height = 30

        self._auto_width(ws)

        # =========================
        # ЛИСТ НАСТРОЕК
        # =========================
        ws2 = wb.create_sheet("Настройки")
        column = 1

        for dataset in datasets:
            name = dataset["name"]
            keywords = dataset["keywords"]
            scenario_name = f"{country} - {name}"

            for keyword in keywords:
                settings = get_settings([keyword])

                ws2.merge_cells(start_row=1, start_column=column, end_row=1, end_column=column + 1)
                cell = ws2.cell(row=1, column=column, value=f"{scenario_name} | {keyword}")
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center")

                ws2.cell(row=2, column=column, value="Настройка")
                ws2.cell(row=2, column=column + 1, value="Значение")

                ws2.cell(row=3, column=column, value="Символ короче")
                ws2.cell(row=4, column=column, value="Символ длиннее")
                ws2.cell(row=5, column=column, value="Номер уже был распознан")
                ws2.cell(row=6, column=column, value="Качество распознавания")
                ws2.cell(row=7, column=column, value="Тип распознавателя")
                ws2.cell(row=8, column=column, value="Режим работы распознавателя")

                if settings:
                    s = settings[0]["params"]

                    ws2.cell(row=3, column=column + 1, value=s.get("min_height_symbol"))
                    ws2.cell(row=4, column=column + 1, value=s.get("max_height_symbol"))
                    ws2.cell(row=5, column=column + 1, value=s.get("same_number_time"))
                    ws2.cell(row=6, column=column + 1, value=s.get("quality"))

                    mode = s.get("mode")
                    if mode == "0":
                        ws2.cell(row=7, column=column + 1, value="Высокая скорость")
                    elif mode == "2":
                        ws2.cell(row=7, column=column + 1, value="Низкая скорость")
                    elif mode == "3":
                        ws2.cell(row=7, column=column + 1, value="Stop & Go")

                    tracking = s.get("tracking_mode")
                    if tracking == "0":
                        ws2.cell(row=8, column=column + 1, value="Дорога/шоссе")
                    elif tracking == "1":
                        ws2.cell(row=8, column=column + 1, value="Парковка")
                    elif tracking == "2":
                        ws2.cell(row=8, column=column + 1, value="Мобильный")
                else:
                    ws2.cell(row=3, column=column + 1, value="Не найдено")

                for row_idx in range(2, 9):
                    for col_idx in [column, column + 1]:
                        cell = ws2.cell(row=row_idx, column=col_idx)
                        cell.border = border
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                column += 3

        self._auto_width(ws2)

        wb.save(output_path)

    # =========================
    # УТИЛИТЫ
    # =========================
    def _auto_width(self, ws):
        for col_idx in range(1, ws.max_column + 1):
            max_length = 0
            col_letter = get_column_letter(col_idx)
            for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
                for cell in row:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except AttributeError:
                        continue
            ws.column_dimensions[col_letter].width = max_length + 4